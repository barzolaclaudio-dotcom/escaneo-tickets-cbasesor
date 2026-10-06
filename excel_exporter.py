import io
import csv
from pathlib import Path
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from ticket_ocr import load_monthly_expenses, get_monthly_summary

def generate_excel_report(month_str: str, user_email: str = None) -> bytes:
    """
    Genera un archivo Excel (.xlsx) profesional con formato institucional CB Asesor:
    - Hoja 1: Resumen Ejecutivo y Totales por Comercio.
    - Hoja 2: Libro IVA Compras (Formato para Contador).
    - Hoja 3: Detalle Ítem por Ítem de Productos Comprados.
    """
    summary = get_monthly_summary(month_str, user_email=user_email)
    tickets = summary.get("tickets", [])
    
    wb = openpyxl.Workbook()
    
    # Estilos de marca CB Asesor
    gold_fill = PatternFill(start_color="D4AF37", end_color="D4AF37", fill_type="solid")
    dark_fill = PatternFill(start_color="121215", end_color="121215", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    gold_font = Font(name="Calibri", size=11, bold=True, color="000000")
    title_font = Font(name="Calibri", size=14, bold=True, color="D4AF37")
    thin_border = Border(
        left=Side(style='thin', color='D4AF37'),
        right=Side(style='thin', color='D4AF37'),
        top=Side(style='thin', color='D4AF37'),
        bottom=Side(style='thin', color='D4AF37')
    )
    
    # -------------------------------------------------------------
    # HOJA 1: Resumen de Gastos por Comercio
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Resumen por Comercio"
    ws1.views.sheetView[0].showGridLines = True
    
    ws1.merge_cells("A1:D1")
    ws1["A1"] = f"CB ASESOR - RESUMEN DE GASTOS ({month_str.replace('_', '/')})"
    ws1["A1"].font = title_font
    
    ws1["A3"] = "TOTAL GASTADO:"; ws1["B3"] = summary["total_spent"]
    ws1["A4"] = "SUBTOTAL NETO:"; ws1["B4"] = summary["total_subtotal"]
    ws1["A5"] = "TOTAL IVA (Crédito Fiscal):"; ws1["B5"] = summary["total_iva"]
    ws1["A6"] = "TOTAL COMPROBANTES:"; ws1["B6"] = summary["tickets_count"]
    
    for row in range(3, 7):
        ws1[f"A{row}"].font = Font(bold=True)
        ws1[f"B{row}"].number_format = '$#,##0.00' if row < 6 else '#,##0'
        
    ws1.append([])
    ws1.append(["Comercio / Proveedor", "Monto Total ($)", "Porcentaje del Gasto", "N° de Compras"])
    header_row_idx = 8
    for col_num in range(1, 5):
        cell = ws1.cell(row=header_row_idx, column=col_num)
        cell.fill = dark_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        
    for v in summary.get("vendor_breakdown", []):
        count_v = sum(1 for t in tickets if t.get("vendor") == v["vendor"])
        ws1.append([v["vendor"], v["total"], f"{v['percentage']}%", count_v])
        
    # -------------------------------------------------------------
    # HOJA 2: Libro IVA Compras (Contador)
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Libro IVA Compras")
    ws2.append(["Fecha", "Tipo Comprobante", "Proveedor", "CUIT", "Neto Gravado ($)", "IVA 21% ($)", "IVA 10.5% ($)", "Total ($)"])
    
    for col_num in range(1, 9):
        cell = ws2.cell(row=1, column=col_num)
        cell.fill = gold_fill
        cell.font = gold_font
        cell.alignment = Alignment(horizontal="center")
        
    for t in tickets:
        ws2.append([
            t.get("date", ""),
            t.get("invoice_type", "Factura A"),
            t.get("vendor", ""),
            t.get("cuit", ""),
            t.get("subtotal", 0.0),
            t.get("iva_21", 0.0),
            t.get("iva_10_5", 0.0),
            t.get("total", 0.0)
        ])
        
    # -------------------------------------------------------------
    # HOJA 3: Detalle Ítem por Ítem de Productos
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="Detalle de Productos")
    ws3.append(["Fecha", "Comercio", "Producto / Ítem", "Cantidad", "Precio ($)"])
    
    for col_num in range(1, 6):
        cell = ws3.cell(row=1, column=col_num)
        cell.fill = dark_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        
    for t in tickets:
        for item in t.get("items", []):
            item_name = item.get("name") or item.get("description") or "Producto / Servicio"
            item_qty = item.get("qty") or item.get("quantity") or 1
            item_price = item.get("price") or item.get("amount") or 0.0
            ws3.append([
                t.get("date", ""),
                t.get("vendor", ""),
                item_name,
                item_qty,
                item_price
            ])
            
    # Ajustar ancho de columnas automáticamente
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)
            
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

def generate_csv_report(month_str: str, user_email: str = None) -> bytes:
    """Genera archivo CSV de IVA Compras listo para enviar al contador."""
    summary = get_monthly_summary(month_str, user_email=user_email)
    tickets = summary.get("tickets", [])
    
    output = io.StringIO()
    writer = csv.writer(output, delimiter=';')
    writer.writerow(["Fecha", "Tipo Comprobante", "Proveedor", "CUIT", "Neto Gravado", "IVA 21%", "IVA 10.5%", "Total"])
    
    for t in tickets:
        writer.writerow([
            t.get("date", ""),
            t.get("invoice_type", "Factura A"),
            t.get("vendor", ""),
            t.get("cuit", ""),
            f"{t.get('subtotal', 0.0):.2f}".replace('.', ','),
            f"{t.get('iva_21', 0.0):.2f}".replace('.', ','),
            f"{t.get('iva_10_5', 0.0):.2f}".replace('.', ','),
            f"{t.get('total', 0.0):.2f}".replace('.', ',')
        ])
        
    return output.getvalue().encode('utf-8-sig')
